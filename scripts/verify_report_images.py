"""Verify the submitted DOCX embeds original image bytes inside the evidence cells."""
import hashlib
import io
from pathlib import Path
from zipfile import ZipFile
from PIL import Image
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT/'모바일웹서비스 프로젝트_공통평가 01_수행 결과 보고서.docx'

def verify():
    document = Document(TARGET)
    assert len(document.tables) == 1
    table = document.tables[0]
    expected_counts = [0,0,1,0,3,1,2,1,1,1,1,1,0,3,2,1,1,1,1,2,1]
    assert len(table.rows) == len(expected_counts)
    images = []
    for index, expected in enumerate(expected_counts):
        cell = table.rows[index].cells[6]
        pictures = cell._tc.xpath('.//wp:inline')
        assert len(pictures) == expected, f'Evidence row {index}: incorrect picture count'
        for picture in pictures:
            name = picture.xpath('./wp:docPr')[0].get('name')
            source = ROOT/'docs/evidence'/name
            original = source.read_bytes()
            blip = picture.xpath('.//a:blip')[0]
            assert blip.get(qn('r:link')) is None, 'Linked images must be embedded'
            stored = document.part.related_parts[blip.get(qn('r:embed'))].blob
            assert hashlib.sha256(original).digest() == hashlib.sha256(stored).digest(), name
            with Image.open(source) as source_image, Image.open(io.BytesIO(stored)) as stored_image:
                assert source_image.size == stored_image.size
                images.append((name, source_image.size))
    assert len(images) == len(document.inline_shapes) == 24
    assert not any(p._p.xpath('.//wp:inline') for p in document.paragraphs), 'Pictures remain outside the table'
    setting = document.settings.element.find(qn('w:doNotAutoCompressPictures'))
    assert setting is not None and setting.get(qn('w:val')) == 'true'
    with ZipFile(TARGET) as package:
        assert package.testzip() is None
    print('PASS: 24 images in their evidence cells; zero pictures below the table.')
    print('PASS: all 24 SHA-256 hashes and pixel dimensions match the original PNGs.')
    print('PASS: all pictures embedded; automatic picture compression disabled.')
    for name, dimensions in images:
        print(f'  {name}: {dimensions[0]}x{dimensions[1]} pixels, byte-identical')

if __name__ == '__main__':
    verify()
