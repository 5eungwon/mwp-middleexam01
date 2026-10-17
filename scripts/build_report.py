"""Fill the supplied report template and embed unmodified screenshot files.
The original requirement table, headings and assessment legend are preserved.
"""
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_BREAK
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/evidence'
TARGET = ROOT / '모바일웹서비스 프로젝트_공통평가 01_수행 결과 보고서.docx'
doc = Document(ROOT / 'docs/templates/report-template.docx')
table = doc.tables[0]
coverage = json.loads((OUT/'coverage.json').read_text())['totals']['percent_covered_display']

# Each requirement refers to full-size evidence in the appendix rather than tiny table thumbnails.
requirements = {
 2: ('완료','○','A01 · ROOT 폴더 캡처'),
 3: ('완료','○','https://github.com/5eungwon/mwp-middleexam01'),
 4: ('완료','○','A03-1 작성 폼\nA03-2 게시 완료\nA03-3 강의 API'),
 5: ('완료','○','A04 · 실제 터미널 TCP 요청/응답'),
 6: ('완료','○','A05-1·2 · 실제 CodeBeautify 검증'),
 7: ('완료','○',f'A06 · Coverage HTML\n17 tests / {coverage}%'),
 8: ('완료','○','A07 · 실제 수신 파일 VS Code 화면'),
 9: ('완료','○','A08 · localhost 브라우저'),
 10: ('완료','○','A09 · Android 시작 화면'),
 11: ('완료','○','A10 · 로컬 강의 API 동기화'),
 12: ('미선택','N/A','UI 제공 방식(상)을 선택함'),
 13: ('완료','○','A11-1·2·3 · 폼/사진 선택/업로드 결과'),
 14: ('완료','○','A12-1·2 · 기존 클라우드 서버/API 동작'),
 15: ('완료','○','A13 · 새 Android와 기존 클라우드 동기화'),
}
for index,(complete,judgment,evidence) in requirements.items():
    row=table.rows[index]
    row.cells[4].text=complete
    row.cells[5].text=judgment
    row.cells[6].text=evidence

extras = [
 ('14-1','즐겨찾기','기기별 즐겨찾기 저장 및 필터. 서버 변경 없이 Client에서 동작.','A14-1'),
 ('14-2','서버 설정 UI','로컬/클라우드 주소 입력과 로그인 UI. 서버 전환 시 기존 인증·캐시 제거.','A14-2'),
 ('14-3','최근 목록 캐시','앱 재실행 시 저장된 목록을 표시. 사진 원본은 네트워크로 수신.','A14-3'),
 ('14-4','내 사진 수정 UI','내 게시물 제목·이야기를 폼으로 수정. 서버의 기존 PATCH API 사용.','A14-4·5'),
 ('14-5','삭제 확인 UI','삭제 전 제목을 표시하고 확인/취소 제공. 서버의 기존 DELETE API 사용.','A14-6'),
]
for index,(ident,name,description,evidence) in enumerate(extras,16):
    row=table.rows[index]
    if index==16:
        row.cells[0].text='14\n(14-1)'
        row.cells[3].text += '\n\n선택 기능 1: '+name+'\n'+description
    else:
        row.cells[0].text=ident
        row.cells[1].text='Client'
        row.cells[2].text=name
        row.cells[3].text=description
    row.cells[4].text='완료';row.cells[5].text='○';row.cells[6].text=evidence
    for cell in row.cells[:4]:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs: run.font.size=Pt(8)

# Preserve table dimensions and reduce only added status text to fit the supplied narrow columns.
for row in table.rows[2:]:
    for cell in row.cells[4:]:
        for p in cell.paragraphs:
            for run in p.runs: run.font.size=Pt(9)

def paragraph(text): doc.add_paragraph(text)

def evidence(code,title,description,filename):
    doc.add_page_break()
    doc.add_heading(f'증빙 {code} · {title}',1)
    paragraph(description)
    image=OUT/filename
    if not image.exists():
        paragraph('미완료: 실제 캡처 파일이 아직 없습니다. '+filename)
        return
    with Image.open(image) as im: w,h=im.size
    section=doc.sections[-1]
    maxw=(section.page_width-section.left_margin-section.right_margin)/914400
    # Only set the display size. python-docx embeds the original pixel bytes without resampling.
    width=min(maxw,7.6*w/h)
    doc.add_picture(str(image),width=Inches(width))
    paragraph(f'원본: docs/evidence/{filename} · {w}×{h}px (원본 해상도 유지)')

doc.add_page_break()
doc.add_heading('구현 및 검증 개요',1)
paragraph('프로젝트: 빛의 기록 · Django PhotoBlogServer / Java Android PhotoViewer')
paragraph('기준 자료: 제공된 ch05_모바일 앱으로 사용하는 기본포토블로그_v1.1.pdf 및 제공 Word 원본 양식. 이 보고서는 원본 표를 유지하고 관련 증빙을 부록에 첨부했습니다. ch04 TDD 강의자료는 제공되지 않았으며, 양식에 명시된 Python 자동 테스트·Code Coverage HTML 요구사항을 구현했습니다.')
paragraph('강의 API: /api_root/Post/ → JSON Array. author(숫자), title, text, created_date, published_date, image를 제공하며 id와 author_name을 추가했습니다. created_date/published_date는 내부 생성/수정 시각에 대응합니다. 확장 API /api/posts/는 페이지형 응답을 유지합니다.')
paragraph(f'검증 결과: Django APITestCase 17개 통과, 분기 포함 Code Coverage {coverage}%. 측정 범위는 posts 앱의 API/웹/모델/폼 코드이며 테스트·마이그레이션·데모 생성 코드는 제외했습니다. HTML 원본은 docs/coverage/index.html입니다.')
paragraph('Android는 강의의 /api_root/Post/를 기본 사용하고 RecyclerView/ImageAdapter로 사진을 표시합니다. deprecated AsyncTask는 ExecutorService와 Handler로 대체했습니다. 사진 선택은 ACTION_OPEN_DOCUMENT를 사용하므로 전체 저장소 접근 권한은 필요하지 않습니다. 네트워크 HTTP 허용은 에뮬레이터/루프백에 제한하고 클라우드는 HTTPS를 사용합니다.')
paragraph('환경: macOS arm64 / Python 3.12 / Django 5.2.18 / DRF 3.16.1 / Android Pixel 7 Android 15(API 35), compileSdk 36. Android assembleDebug와 lintDebug를 실행했습니다.')
paragraph('PythonAnywhere: 기존 배포 서버 https://asdfgh0986.pythonanywhere.com/api_root/Post/ 가 HTTP 200과 강의 JSON Array를 반환합니다. 새 Android 앱이 이 서버의 실제 게시물과 사진을 수신하는 호환 연동을 검증했습니다. 요구사항 12·13은 이 기존 클라우드 서버와 새 클라이언트의 실제 동작을 증빙합니다. 저장소의 새 Django 소스를 PythonAnywhere에 재배포한 것은 아니며, 로컬 서버와 기존 클라우드 서버의 UI·확장 기능은 다릅니다.')
paragraph('작성일: '+datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d')+' · 학번/이름은 제출자가 첫 페이지에 기입해야 합니다.')

entries = [
 ('A01','폴더 구조','제공된 ROOT 구조의 서버·클라이언트·Word 보고서가 보이는 실제 Finder 화면.','01-folder-structure.png'),
 ('A03-1','POST 작성 완료 화면','강의 7쪽: 제목·설명·사진 선택이 완료된 실제 웹 입력 폼. URL 주소창 포함.','03-web-post-form.png'),
 ('A03-2','POST 게시 완료 화면','폼을 실제 제출한 후 저장된 게시물 상세 페이지. URL 주소창 포함.','03-web-post-detail.png'),
 ('A03-3','강의 REST API','강의 5~7쪽: /api_root/Post/ JSON Array API의 실제 브라우저 화면.','03-lecture-api-browser.png'),
 ('A04','API 검증 (2/3)','강의 9쪽: Python TCP socket으로 Telnet과 같은 원시 HTTP 요청을 전송했습니다. 실제 터미널에서 Request/Response line, headers, 빈 줄, JSON Array 및 PASS 결과를 확인했습니다. 원본은 04-tcp-response.txt와 04-api-array.json입니다.','04-terminal.png'),
 ('A05-1','API 검증 (3/3) CodeBeautify','강의 10쪽: 실제 CodeBeautify JSON Viewer에 서버의 공개 JSON 배열을 입력하고 Beautify와 Validate를 실행한 화면. 토큰·비밀번호는 입력하지 않았습니다.','05-codebeautify-code.png'),
 ('A05-2','CodeBeautify Tree 확인','같은 실제 API 응답을 Tree 모드로 펼쳐 배열과 필드 구조를 확인했습니다.','05-codebeautify-tree.png'),
 ('A06','Code Coverage HTML','제공 양식 요구사항 6: coverage run --branch / coverage html로 생성한 실제 HTML 보고서. 원본 보고서와 테스트 로그를 저장소에 포함했습니다.','06-coverage-html.png'),
 ('A07','Socket Server 수신 메시지','강의 24쪽: 앱을 8001 포트의 TCP 수신 프록시에 연결했습니다. 실제 GET /api_root/Post/와 Authorization/User-Agent/Host 등의 수신 헤더를 파일로 저장해 VS Code에서 열었습니다. 토큰과 Cookie는 파일 저장 전 가렸고 요청 본문·비밀번호는 저장하지 않았습니다. 프록시는 요청을 로컬 Django :8000에 전달합니다.','07-socket-vscode.png'),
 ('A08','Server 브라우저 출력','localhost Django 서버 갤러리의 실제 브라우저 화면.','08-server-browser.png'),
 ('A09','Client 시작 화면','강의 25쪽: 강의 API가 기본인 새 Android 앱의 시작 화면.','09-client-start.png'),
 ('A10','Client 로컬 동기화 완료','강의 25쪽: /api_root/Post/ JSON Array를 수신하고 사진 목록을 표시한 실제 에뮬레이터 화면.','10-client-lecture-synced.png'),
 ('A11-1','Upload UI (상)','강의 19쪽: onClickUpload → openEditor. 제목·이야기를 직접 입력하는 UI 방식 선택.','11-upload-form.png'),
 ('A11-2','Upload 사진 선택 완료','시스템 사진 선택기로 사진을 선택한 실제 화면.','11-upload-selected.png'),
 ('A11-3','Upload 서버 저장 및 동기화','uploadPost에서 multipart POST, 토큰 인증, 성공 후 강의 API 재동기화를 수행한 실제 완료 화면.','11-upload-complete.png'),
 ('A12-1','PythonAnywhere 서버 브라우저','사용자 계정의 기존 배포 블로그가 실제 HTTPS 주소에서 동작하는 화면. 새 Django 소스를 재배포했다는 증거는 아닙니다.','12-cloud-server.png'),
 ('A12-2','PythonAnywhere 강의 API','기존 서버의 정확한 /api_root/Post/ 경로가 JSON Array를 응답하는 실제 주소창 포함 화면.','12-cloud-api.png'),
 ('A13','Client 클라우드 동기화','새 Android 앱의 서버 주소를 기존 PythonAnywhere HTTPS 주소로 설정하고 실제 게시물과 꽃 사진을 수신한 화면.','13-cloud-client-synced.png'),
 ('A14-1','즐겨찾기','Client 로컬 즐겨찾기 저장 및 필터 화면.','14-favorite.png'),
 ('A14-2','서버 설정','Client 서버 주소 및 강의 API 모드 설정 화면.','14-settings.png'),
 ('A14-3','최근 목록 캐시','앱 재실행 후 캐시된 목록과 즐겨찾기 유지 확인. 사진 원본은 네트워크 수신.','14-cache.png'),
 ('A14-4','수정 입력 폼','Client 본인 게시물 제목·이야기 수정 UI.','14-edit-form.png'),
 ('A14-5','수정 완료','수정 결과를 서버와 동기화한 실제 화면.','14-edit-complete.png'),
 ('A14-6','삭제 확인','제목을 표시하는 삭제 확인창. 캡처 후 취소했으며 실제 DELETE와 권한 검증은 자동 테스트로 확인.','14-delete-confirm.png'),
]
for entry in entries: evidence(*entry)
doc.add_page_break();doc.add_heading('검증 범위 및 재현 방법',1)
paragraph('클라우드는 기존 배포 서버와 새 Android의 읽기/동기화 호환성을 검증했습니다. 새 Django 소스로 클라우드의 웹 화면·로그인·업로드 기능까지 교체하려면 계정 콘솔에서 docs/PYTHONANYWHERE.md의 WSGI·static/media 전환 절차를 추가로 수행해야 합니다. 보고서의 클라우드 완료 판정은 기존 서버 동작과 새 클라이언트 동기화에 한정됩니다.')
paragraph('테스트: cd PhotoBlogServer → python -m coverage run manage.py test -v 2 → python -m coverage html. Android: cd PhotoViewer → ./gradlew :app:assembleDebug :app:lintDebug. 상세 실행·인증·캡처 재현 방법은 README.md를 참조합니다.')
paragraph('작성된 코드: posts/serializers.py, posts/views.py, posts/forms.py, posts/tests.py, MainActivity.java, ImageAdapter.java, scripts/socket_capture.py, scripts/verify_raw_http.py. 데모 풍경은 Pillow로 직접 생성한 일러스트입니다.')
doc.save(TARGET)
missing=[name for *_,name in entries if not (OUT/name).exists()]
print(TARGET.name,'| embedded screenshots:',len(doc.inline_shapes),'| missing:',missing)
